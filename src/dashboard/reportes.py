"""
Utilidades para generar reportes DOCX y para buscar literatura relevante.

La generación de DOCX (generar_reporte_docx) formatea datos reales que ya
tiene el sistema -- sin riesgo de inventar nada, es solo maquetación.

La búsqueda de literatura consulta en vivo la API pública de NCBI
(PubMed E-utilities) dados los marcadores utilizados y el diagnóstico de 
NeuroInsight. Si la búsqueda no encuentra nada o falla la conexión, 
devuelve una lista vacía.
"""

from datetime import datetime
import io

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import json
import urllib.parse
import urllib.request


def generar_reporte_docx(
    ficha_paciente: dict = None,
    probabilidades: dict = None,
    clase_predicha: str = None,
    factores_riesgo: list = None,
    papers: list = None,
    notas: str = None,
    incluir_prediccion: bool = True,
    incluir_literatura: bool = False,
    incluir_notas: bool = False,
) -> io.BytesIO:
    """
    Genera un reporte DOCX con las secciones que se soliciten -- pensado
    para la página "Generar informe" (5.3), donde el usuario elige qué
    incluir en vez de recibir siempre el mismo documento fijo.

    Args:
        ficha_paciente: Diccionario con datos del paciente (solo se usa si incluir_prediccion)
        probabilidades: Dict con probabilidades por clase {clase: probabilidad}
        clase_predicha: Clase predicha (CN, MCI, AD)
        factores_riesgo: Lista de factores de riesgo identificados (opcional)
        papers: Lista de papers (de obtener_papers_relevantes) -- solo se usa si incluir_literatura
        notas: Texto de notas clínicas -- solo se usa si incluir_notas
        incluir_prediccion / incluir_literatura / incluir_notas: qué secciones añadir
    """
    
    doc = Document()
    
    # Encabezado
    title = doc.add_heading('NeuroInsight', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_heading('Informe del paciente', level=2)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    fecha = datetime.now().strftime("%d/%m/%Y %H:%M")
    fecha_p = doc.add_paragraph(f"Generado: {fecha}")
    fecha_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fecha_p.runs[0].font.size = Pt(10)
    fecha_p.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    
    doc.add_paragraph()  # Salto
    
    # Disclaimer
    disclaimer = doc.add_paragraph()
    disclaimer.add_run("⚠️ DISCLAIMER - IMPORTANTE\n").font.bold = True
    disclaimer.add_run(
        "Este reporte es una herramienta de APOYO CLÍNICO experimental y NO constituye un diagnóstico médico. "
        "Las predicciones están basadas en modelos de machine learning entrenados con datos del estudio ADNI. "
        "Cualquier decisión clínica debe ser tomada por profesionales médicos calificados tras evaluación clínica completa. "
        "Esta herramienta es complementaria, no sustituye el juicio clínico."
    )
    disclaimer_format = disclaimer.paragraph_format
    disclaimer_format.left_indent = Inches(0.3)
    disclaimer.style = 'List Bullet'
    
    doc.add_paragraph()

    numero_seccion = 1

    # Sección: Predicción (datos + resultado)
    if incluir_prediccion and ficha_paciente is not None:
        doc.add_heading(f'{numero_seccion}. Datos del Paciente', level=1)
        numero_seccion += 1

        datos_table = doc.add_table(rows=1, cols=2)
        datos_table.style = 'Light Grid Accent 1'
        hdr_cells = datos_table.rows[0].cells
        hdr_cells[0].text = 'Variable'
        hdr_cells[1].text = 'Valor'

        for key, valor in ficha_paciente.items():
            if valor is not None:
                row_cells = datos_table.add_row().cells
                row_cells[0].text = key.replace('_', ' ').title()
                row_cells[1].text = str(valor)

        doc.add_paragraph()

        if clase_predicha:
            doc.add_heading(f'{numero_seccion}. Resultado de la Predicción', level=1)
            numero_seccion += 1

            pred_text = doc.add_paragraph()
            pred_text.add_run("Diagnóstico predicho: ").bold = True
            pred_text.add_run(f"{clase_predicha}\n")

            diagnosticos = {
                "CN": "Cognitivamente Normal (sin deterioro)",
                "MCI": "Deterioro Cognitivo Leve (MCI - mild cognitive impairment)",
                "AD": "Enfermedad de Alzheimer",
            }
            pred_text.add_run(f"Significado: {diagnosticos.get(clase_predicha, 'Desconocido')}")

            doc.add_paragraph()

            if probabilidades:
                prob_table = doc.add_table(rows=1, cols=2)
                prob_table.style = 'Light Grid Accent 1'
                hdr_cells = prob_table.rows[0].cells
                hdr_cells[0].text = 'Categoría'
                hdr_cells[1].text = 'Probabilidad'
                for clase, prob in probabilidades.items():
                    row_cells = prob_table.add_row().cells
                    row_cells[0].text = clase
                    row_cells[1].text = f"{prob*100:.1f}%"

            doc.add_paragraph()

    # Sección: "Por qué" -- factores que más pesaron (SHAP), independiente de incluir_prediccion
    # para poder incluirse aunque no se incluyan los datos crudos del paciente.
    if factores_riesgo:
        doc.add_heading(f'{numero_seccion}. Por Qué (Factores Determinantes)', level=1)
        numero_seccion += 1
        doc.add_paragraph(
            "Variables ordenadas por peso en esta predicción concreta (método SHAP). "
            "Un peso mayor indica más influencia en el resultado."
        )
        for factor in factores_riesgo:
            doc.add_paragraph(factor, style='List Bullet')
        doc.add_paragraph()

    # Sección: Literatura relevante (papers reales de PubMed, ver obtener_papers_relevantes)
    if incluir_literatura:
        doc.add_heading(f'{numero_seccion}. Literatura Relevante (PubMed)', level=1)
        numero_seccion += 1
        if papers:
            for paper in papers:
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(f"{paper.get('titulo', '(sin título)')} ").bold = True
                autores = ", ".join(paper.get("autores", [])[:3])
                p.add_run(f"-- {autores} ({paper.get('año', '----')}), {paper.get('journal', '')}. ")
                p.add_run(f"PMID {paper.get('pmid', '')} -- https://pubmed.ncbi.nlm.nih.gov/{paper.get('pmid', '')}/")
        else:
            doc.add_paragraph("No se encontraron artículos relevantes en PubMed para este perfil en el momento de generar el informe.")
        doc.add_paragraph()

    # Sección: Notas clínicas
    if incluir_notas:
        doc.add_heading(f'{numero_seccion}. Notas Clínicas', level=1)
        numero_seccion += 1
        if notas and notas.strip():
            doc.add_paragraph(notas)
        else:
            doc.add_paragraph("(Sin notas registradas para este paciente)")
        doc.add_paragraph()

    # Footer
    footer_text = doc.add_paragraph(
        "---\nEste reporte fue generado automáticamente por NeuroInsight. "
        "Es una herramienta de investigación / TFM. Consultar documentación técnica para detalles del modelo."
    )
    footer_text.runs[0].font.size = Pt(9)
    footer_text.runs[0].font.color.rgb = RGBColor(150, 150, 150)
    
    # Guardar en BytesIO
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    
    return output

# ---------------------------------------------------------------------
# Búsqueda de literatura relevante en PubMed (API real de NCBI)
# ---------------------------------------------------------------------

BASE_PUBMED = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils"

TERMINOS_BIOMARCADOR = {
    "tau_pg_ml": "tau protein",
    "p_tau_pg_ml": "phosphorylated tau",
    "abeta42_pg_ml": "amyloid beta 42",
    "ptau217_plasma_pg_ml": "plasma p-tau217",
    "ratio_ab42_ab40_plasma": "amyloid beta 42/40 ratio plasma",
    "hipocampo_mm3": "hippocampal volume",
    "mmse": "MMSE cognitive assessment",
    "apoe4_positivo": "APOE4",
    # Variantes ya formateadas (title-case, con espacios) que llegan desde prediccion.py
    "Tau Pg Ml": "tau protein",
    "Abeta42 Pg Ml": "amyloid beta 42",
    "Ptau217 Plasma Pg Ml": "plasma p-tau217",
    "Ratio Ab42 Ab40 Plasma": "amyloid beta 42/40 ratio plasma",
}

TERMINOS_DIAGNOSTICO = {
    "AD": "Alzheimer disease",
    "MCI": "mild cognitive impairment",
    "CN": "cognitively normal aging",
}


def _get_json(ruta: str, params: dict) -> dict:
    url = f"{BASE_PUBMED}/{ruta}?" + urllib.parse.urlencode(params)
    with urllib.request.urlopen(url, timeout=10) as resp:
        return json.loads(resp.read().decode("utf-8"))


def obtener_papers_relevantes(biomarcadores: list, diagnostico: str, max_resultados: int = 5) -> list:
    """Busca en PubMed papers relacionados con los biomarcadores dados y
    el diagnóstico -- API real de NCBI, no una lista fija. Cada resultado
    es verificable en https://pubmed.ncbi.nlm.nih.gov/{pmid}/ porque el
    pmid procede directamente de esa búsqueda.

    Devuelve lista vacía si no hay resultados o falla la conexión --
    nunca un resultado inventado como sustituto.
    """
    if not biomarcadores:
        biomarcadores = []
    if isinstance(biomarcadores, str):
        biomarcadores = [biomarcadores]

    termino_dx = TERMINOS_DIAGNOSTICO.get(diagnostico, diagnostico)
    if biomarcadores:
        terminos_bio = [TERMINOS_BIOMARCADOR.get(b, b) for b in biomarcadores]
        bloque_bio = " OR ".join(f'"{t}"[Title/Abstract]' for t in terminos_bio)
        consulta = f'({bloque_bio}) AND "{termino_dx}"[Title/Abstract]'
    else:
        consulta = f'"{termino_dx}"[Title/Abstract]'
    parametros_comunes = {"tool": "neuroinsight-tfm", "email": "no-reply@example.com"}

    try:
        busqueda = _get_json("esearch.fcgi", {
            "db": "pubmed", "term": consulta, "retmax": max_resultados,
            "sort": "relevance", "retmode": "json", **parametros_comunes,
        })
        ids = busqueda.get("esearchresult", {}).get("idlist", [])
        if not ids:
            return []

        resumenes = _get_json("esummary.fcgi", {
            "db": "pubmed", "id": ",".join(ids), "retmode": "json", **parametros_comunes,
        })
        resultado = resumenes.get("result", {})

        papers = []
        for pmid in ids:
            item = resultado.get(pmid)
            if not item:
                continue
            autores = [a.get("name", "") for a in item.get("authors", []) if a.get("name")]
            papers.append({
                "titulo": item.get("title", "(sin título)").rstrip("."),
                "autores": autores or ["(autores no disponibles)"],
                "año": (item.get("pubdate", "") or "----")[:4],
                "journal": item.get("fulljournalname") or item.get("source", ""),
                "pmid": pmid,
            })
        return papers
    except Exception:
        return []